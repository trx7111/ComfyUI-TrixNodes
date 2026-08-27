import os
import re
import json
import io
import time
import logging
import base64
from typing import List, Dict, Any, Optional
from PIL import Image

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.drawing.image import Image as OpenPyXLImage
import xlsxwriter

from .excel_manager import ExcelStyleManager

logger = logging.getLogger("ComfyUI-TrixPromptAIO.StyleManager")


class StyleManager:
    """Manager for prompt styles, presets, categories, and favorites directly using Excel files."""

    INCREMENT_COUNTERS: Dict[str, int] = {}
    LAST_CHOSEN_INDEX: Dict[str, int] = {}
    LAST_GROUP_SEED: Dict[str, int] = {}
    LAST_EXPANDED_DETAILS: Dict[str, List[dict]] = {}
    NODE_INCREMENT_SEEDS: Dict[str, int] = {}
    LAST_START_SEED: Dict[str, int] = {}
    FAVORITES_SET: set = set()

    @staticmethod
    def _clean_obsolete_files():
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        db_json = os.path.join(base_dir, "styles_db.json")
        if os.path.exists(db_json):
            try:
                os.remove(db_json)
            except Exception:
                pass
        
        thumbs_dir = os.path.join(base_dir, "styles", "thumbnails")
        if os.path.exists(thumbs_dir):
            try:
                import shutil
                shutil.rmtree(thumbs_dir, ignore_errors=True)
            except Exception:
                pass

    @staticmethod
    def load_styles() -> List[Dict[str, Any]]:
        StyleManager._clean_obsolete_files()
        styles, _ = ExcelStyleManager.load_master_data()
        for s in styles:
            s_id = s.get("id")
            if s_id and s_id in StyleManager.FAVORITES_SET:
                s["favorite"] = True
        return styles

    @staticmethod
    def toggle_favorite(style_id: str) -> List[Dict[str, Any]]:
        if style_id in StyleManager.FAVORITES_SET:
            StyleManager.FAVORITES_SET.remove(style_id)
        else:
            StyleManager.FAVORITES_SET.add(style_id)
        return StyleManager.load_styles()

    @staticmethod
    def _write_category_to_disk(cat_clean: str, all_styles: List[Dict[str, Any]], all_images: Dict[str, tuple]):
        """Re-writes the entire category .xlsx file from RAM using xlsxwriter for perfect compatibility."""
        styles_dir = ExcelStyleManager.get_styles_dir()
        cat_path = os.path.join(styles_dir, f"{cat_clean}.xlsx")
        tmp_path = cat_path + '.tmp'
        
        styles_for_cat = [s for s in all_styles if s.get('categoryKey') == cat_clean or s.get('category') == cat_clean]
        
        wb = xlsxwriter.Workbook(tmp_path)
        ws = wb.add_worksheet(cat_clean)

        cell_format = wb.add_format({'text_wrap': True, 'valign': 'top'})
        
        # Set columns with format
        ws.set_column('A:B', 16, cell_format)
        ws.set_column('C:C', 24, cell_format)
        ws.set_column('D:D', 20, cell_format)
        ws.set_column('E:E', 55, cell_format)
        ws.set_column('F:F', 35, cell_format)

        headers = ["Preview 1", "Preview 2", "Name", "Tag", "Positive Prompt", "Negative Prompt"]
        data = []
        for s in styles_for_cat:
            data.append(["", "", s.get("name", ""), s.get("tag", ""), s.get("positive", ""), s.get("negative", "")])

        end_row = len(data) + 1 if len(data) > 0 else 2
        table_data = data if data else [["", "", "", "", "", ""]] # Table must have at least one row
        
        ws.add_table(f'A1:F{end_row}', {
            'data': table_data,
            'columns': [{'header': h} for h in headers],
            'style': 'Table Style Medium 2'
        })

        ws.set_default_row(hide_unused_rows=False)

        for row_idx, s in enumerate(styles_for_cat, start=1):
            ws.set_row(row_idx, 78) # Height
            thumb_val = s.get("thumbnail")
            if not thumb_val:
                continue
            
            t1 = thumb_val[0] if isinstance(thumb_val, list) and len(thumb_val) > 0 else (thumb_val if isinstance(thumb_val, str) else None)
            t2 = thumb_val[1] if isinstance(thumb_val, list) and len(thumb_val) > 1 else None

            if t1 and t1 in all_images:
                img_bytes, mime = all_images[t1]
                try:
                    with Image.open(io.BytesIO(img_bytes)) as pil_img:
                        w, h = pil_img.size
                        scale = 88.0 / max(1, max(w, h))
                    ws.insert_image(row_idx, 0, t1, {'image_data': io.BytesIO(img_bytes), 'x_scale': scale, 'y_scale': scale, 'x_offset': 4, 'y_offset': 4, 'object_position': 1})
                except Exception as e:
                    logger.error(f"Error scaling image {t1}: {e}")
            
            if t2 and t2 in all_images:
                img_bytes, mime = all_images[t2]
                try:
                    with Image.open(io.BytesIO(img_bytes)) as pil_img:
                        w, h = pil_img.size
                        scale = 88.0 / max(1, max(w, h))
                    ws.insert_image(row_idx, 1, t2, {'image_data': io.BytesIO(img_bytes), 'x_scale': scale, 'y_scale': scale, 'x_offset': 4, 'y_offset': 4, 'object_position': 1})
                except Exception as e:
                    logger.error(f"Error scaling image {t2}: {e}")

        wb.close()
        
        # Replace original safely
        try:
            if os.path.exists(cat_path):
                os.remove(cat_path)
            os.rename(tmp_path, cat_path)
            logger.info(f"Rebuilt {cat_clean}.xlsx with {len(styles_for_cat)} styles using xlsxwriter.")
        except Exception as err:
            logger.error(f"Error replacing {cat_clean}.xlsx: {err}")
        finally:
            if os.path.exists(tmp_path):
                try: os.remove(tmp_path)
                except Exception: pass

    @staticmethod
    def create_category(category_key: str) -> List[Dict[str, Any]]:
        cat_clean = ExcelStyleManager._sanitize_name(category_key)
        styles_dir = ExcelStyleManager.get_styles_dir()
        cat_path = os.path.join(styles_dir, f"{cat_clean}.xlsx")

        if not os.path.exists(cat_path):
            StyleManager._write_category_to_disk(cat_clean, [], {})

        ExcelStyleManager.clear_cache()
        return StyleManager.load_styles()

    @staticmethod
    def _prepare_openpyxl_img(base64_str: str) -> Optional[OpenPyXLImage]:
        if not base64_str or not isinstance(base64_str, str):
            return None
        try:
            data = base64_str
            if "," in data:
                data = data.split(",", 1)[1]
            img_bytes = base64.b64decode(data)
            with Image.open(io.BytesIO(img_bytes)) as img:
                # Always convert to JPEG for maximum xlsx compatibility
                # (PNG/WEBP with UUID rIds causes rendering failures in PlanMaker, mobile apps)
                if img.mode in ("RGBA", "P", "LA"):
                    bg = Image.new("RGB", img.size, (255, 255, 255))
                    if img.mode == "P":
                        img = img.convert("RGBA")
                    if img.mode in ("RGBA", "LA"):
                        bg.paste(img, mask=img.split()[-1])
                    else:
                        bg.paste(img)
                    img = bg
                elif img.mode != "RGB":
                    img = img.convert("RGB")

                out_buf = io.BytesIO()
                img.save(out_buf, format="JPEG", quality=90, optimize=True)
                out_buf.seek(0)

                xl_img = OpenPyXLImage(out_buf)
                xl_img.width = 96
                xl_img.height = 96
                return xl_img
        except Exception as e:
            logger.error(f"Error processing base64 image: {e}")
            return None

    @staticmethod
    def save_custom_style(style_data: Dict[str, Any]) -> List[Dict[str, Any]]:
        name = str(style_data.get("name") or "").strip() or "Untitled Style"
        category = style_data.get("categoryKey") or style_data.get("category") or "custom"
        cat_clean = ExcelStyleManager._sanitize_name(category)
        
        positive = str(style_data.get("positive") or style_data.get("prompt") or "").strip()
        negative = str(style_data.get("negative") or "").strip()
        
        tag_val = str(style_data.get("tag") or "").strip()
        if not tag_val:
            tag_val = "@" + re.sub(r'[^a-z0-9]+', '_', name.lower()).strip('_')

        # Load current data
        all_styles, all_images = ExcelStyleManager.load_master_data()

        # Find if style exists
        search_id = style_data.get("id") or name
        target_style = next((s for s in all_styles if s.get("id") == search_id or s.get("name") == name), None)
        
        modified_cats = {cat_clean}
        if target_style:
            old_cat = ExcelStyleManager._sanitize_name(target_style.get("categoryKey") or target_style.get("category") or "custom")
            if old_cat and old_cat != cat_clean:
                modified_cats.add(old_cat)
        else:
            target_style = {}
            all_styles.append(target_style)

        target_style["name"] = name
        target_style["categoryKey"] = cat_clean
        target_style["category"] = cat_clean
        target_style["tag"] = tag_val
        target_style["positive"] = positive
        target_style["negative"] = negative

        img1_b64 = style_data.get("image1_base64")
        img2_b64 = style_data.get("image2_base64")

        # Process new images and add to RAM cache
        thumb_list = []
        
        def process_b64(b64_str, existing_key):
            if b64_str:
                try:
                    data = b64_str.split(",", 1)[1] if "," in b64_str else b64_str
                    img_bytes = base64.b64decode(data)
                    with Image.open(io.BytesIO(img_bytes)) as img:
                        if img.mode in ("RGBA", "P", "LA"):
                            bg = Image.new("RGB", img.size, (255, 255, 255))
                            if img.mode == "P":
                                img = img.convert("RGBA")
                            if img.mode in ("RGBA", "LA"):
                                bg.paste(img, mask=img.split()[-1])
                            else:
                                bg.paste(img)
                            img = bg
                        elif img.mode != "RGB":
                            img = img.convert("RGB")
                        
                        out_buf = io.BytesIO()
                        img.save(out_buf, format="JPEG", quality=90, optimize=True)
                        final_bytes = out_buf.getvalue()
                        
                        # Generate a unique key for the new image
                        new_key = f"{cat_clean}_{name}_{int(time.time()*1000)}.jpeg"
                        new_key = re.sub(r'[^a-zA-Z0-9_\-\.]', '_', new_key)
                        all_images[new_key] = (final_bytes, "image/jpeg")
                        return new_key
                except Exception as e:
                    logger.error(f"Error processing base64 image: {e}")
            return existing_key

        existing_thumb = target_style.get("thumbnail") or style_data.get("thumbnail")
        t1 = existing_thumb[0] if isinstance(existing_thumb, list) and len(existing_thumb) > 0 else (existing_thumb if isinstance(existing_thumb, str) else None)
        t2 = existing_thumb[1] if isinstance(existing_thumb, list) and len(existing_thumb) > 1 else None

        new_t1 = process_b64(img1_b64, t1)
        new_t2 = process_b64(img2_b64, t2)

        if new_t1: thumb_list.append(new_t1)
        if new_t2: thumb_list.append(new_t2)
        
        target_style["thumbnail"] = thumb_list

        # Save all modified categories to disk
        for cat in modified_cats:
            StyleManager._write_category_to_disk(cat, all_styles, all_images)

        ExcelStyleManager.clear_cache()
        return StyleManager.load_styles()

    @staticmethod
    def delete_style(style_id: str) -> List[Dict[str, Any]]:
        all_styles, all_images = ExcelStyleManager.load_master_data()
        target_style = next((s for s in all_styles if s.get("id") == style_id or s.get("name") == style_id), None)
        if not target_style:
            return all_styles

        cat_clean = target_style.get("categoryKey") or target_style.get("category") or "custom"
        
        # Remove from RAM
        all_styles.remove(target_style)

        # Write the updated category back to disk
        StyleManager._write_category_to_disk(cat_clean, all_styles, all_images)

        ExcelStyleManager.clear_cache()
        return StyleManager.load_styles()

    @staticmethod
    def move_styles_category(style_ids: List[str], target_category: str) -> List[Dict[str, Any]]:
        target_cat_clean = ExcelStyleManager._sanitize_name(target_category)
        
        all_styles, all_images = ExcelStyleManager.load_master_data()
        
        # Keep track of which categories were modified so we can rewrite them
        modified_cats = {target_cat_clean}
        
        for s in all_styles:
            if s.get("id") in style_ids or s.get("name") in style_ids:
                old_cat = s.get("categoryKey") or s.get("category")
                if old_cat and old_cat != target_cat_clean:
                    modified_cats.add(old_cat)
                
                s["categoryKey"] = target_cat_clean
                s["category"] = target_cat_clean
        
        for cat in modified_cats:
            StyleManager._write_category_to_disk(cat, all_styles, all_images)
            
        ExcelStyleManager.clear_cache()
        return StyleManager.load_styles()

    @staticmethod
    def delete_category(category_key: str, delete_styles: bool = False) -> List[Dict[str, Any]]:
        cat_clean = ExcelStyleManager._sanitize_name(category_key)
        cat_path = os.path.join(ExcelStyleManager.get_styles_dir(), f"{cat_clean}.xlsx")

        if os.path.exists(cat_path):
            if delete_styles:
                os.remove(cat_path)
            else:
                all_styles = StyleManager.load_styles()
                cat_styles = [s.get("id") or s.get("name") for s in all_styles if (s.get("categoryKey") == cat_clean or s.get("category") == cat_clean) and (s.get("id") or s.get("name"))]
                if cat_styles:
                    StyleManager.move_styles_category(cat_styles, "custom")
                if os.path.exists(cat_path):
                    os.remove(cat_path)

        ExcelStyleManager.clear_cache()
        return StyleManager.load_styles()

    @staticmethod
    def reorder_styles(category_key: str, ordered_ids: List[str]) -> List[Dict[str, Any]]:
        cat_clean = ExcelStyleManager._sanitize_name(category_key)
        all_styles, all_images = ExcelStyleManager.load_master_data()
        
        cat_styles = [s for s in all_styles if s.get("categoryKey") == cat_clean or s.get("category") == cat_clean]
        if not cat_styles:
            return all_styles

        style_map = {}
        for s in cat_styles:
            s_id = s.get("id")
            if s_id:
                style_map[str(s_id)] = s
            s_name = s.get("name")
            if s_name and str(s_name) not in style_map:
                style_map[str(s_name)] = s

        reordered_cat_styles = []
        added_keys = set()
        for s_id in ordered_ids:
            s_str = str(s_id)
            if s_str in style_map:
                item = style_map[s_str]
                if id(item) not in added_keys:
                    reordered_cat_styles.append(item)
                    added_keys.add(id(item))

        for s in cat_styles:
            if id(s) not in added_keys:
                reordered_cat_styles.append(s)
                added_keys.add(id(s))

        StyleManager._write_category_to_disk(cat_clean, reordered_cat_styles, all_images)
        ExcelStyleManager.clear_cache()
        return StyleManager.load_styles()

    @staticmethod
    def expand_tags(text: str, seed: int = 0, group_mode: str = "increment", group_seed: int = 1, node_id: str = "default") -> str:
        if not text or ("@" not in text and "*" not in text):
            return text
        styles = StyleManager.load_styles()
        tag_map = {}
        category_map = {}

        StyleManager.LAST_EXPANDED_DETAILS[node_id] = []

        for s in styles:
            name = s.get("name", "")
            explicit_tag = s.get("tag")
            cat = s.get("categoryKey") or s.get("category") or "custom"
            norm_cat = "*" + re.sub(r'[^a-z0-9]+', '_', cat.strip().lower()).strip('_')
            raw_cat = "*" + cat.strip().lower()
            pos = (s.get("positive") or s.get("prompt") or "").strip()
            
            if name:
                tag = "@" + re.sub(r'[^a-z0-9]+', '_', name.strip().lower()).strip('_')
                tag_map[tag.lower()] = pos
            if explicit_tag:
                t = explicit_tag.strip()
                if not t.startswith("@"):
                    t = "@" + t
                tag_map[t.lower()] = pos
            
            for c_key in set([norm_cat, raw_cat]):
                if c_key not in category_map:
                    category_map[c_key] = []
                category_map[c_key].append({"name": name, "positive": pos})

        effective_seed = max(1, int(group_seed))

        def replacer(match):
            m = match.group(0)
            if m.startswith("@"):
                return tag_map.get(m.lower(), m)
            elif m.startswith("*"):
                norm_cat = "*" + re.sub(r'[^a-z0-9]+', '_', m[1:].strip().lower()).strip('_')
                options = category_map.get(m) or category_map.get(norm_cat) or category_map.get(m.lower())
                if options:
                    key = f"{node_id}_{norm_cat}"
                    N = len(options)
                    idx = (effective_seed - 1) % N

                    StyleManager.LAST_CHOSEN_INDEX[key] = idx
                    selected_item = options[idx]

                    details = StyleManager.LAST_EXPANDED_DETAILS.setdefault(node_id, [])
                    if not any(d.get("tag") == m for d in details):
                        details.append({
                            "tag": m,
                            "name": selected_item.get("name") or "Style",
                            "index": idx,
                            "total": N
                        })

                    return selected_item["positive"]
                return m
            return m

        current_text = text
        for _ in range(3):
            next_text = re.sub(r'[@\*][a-zA-Z0-9_-]+', replacer, current_text)
            if next_text == current_text:
                break
            current_text = next_text

        return current_text
